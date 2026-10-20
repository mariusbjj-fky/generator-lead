"""Excel output (+ CSV of the merged list) and the end-of-run summary."""
import csv
import os

from . import core as C

# Excel caps sheet names at 31 characters, so the merged tab is shortened.
TABS = [
    ("Cat 1 - New Businesses", "1"),
    ("Cat 2A - Inactive No Ads", "2A"),
    ("Cat 2B - Inactive With Ads", "2B"),
    ("Cat 3 - No Website", "3"),
]
ALL_TAB = "All Leads (dedup by CUI)"

OUT = [  # (header, function(row) -> value)
    ("Lead ID", lambda r: r["lead_id"]),
    ("Priority", lambda r: r["priority"]),
    ("Score", lambda r: int(r["score"]) if str(r["score"]).isdigit() else ""),
    ("Categories", lambda r: r["categories"]),
    ("Legal name", lambda r: r["legal_name"]),
    ("Brand name", lambda r: r["brand_name"]),
    ("Industry", lambda r: C.cfg()["industries"].get(r["industry"], {}).get("label", r["industry"])),
    ("CAEN code", lambda r: r["caen_code"]),
    ("CAEN version", lambda r: r["caen_version"]),
    ("City", lambda r: r["city"]),
    ("County (județ)", lambda r: r["county"]),
    ("Registration date", lambda r: r["registration_date"]),
    ("CUI/CIF", lambda r: r["cui"]),
    ("Registry verification", lambda r: (r["verification_status"] + (f" ({r['verified_by']})" if r["verified_by"] else ""))
     if r["verification_status"] else ""),
    ("Website URL", lambda r: "NONE" if r["website_status"] == "NONE" else (r["website_url"] or (
        "NOT FOUND" if r["website_status"] == "NOT_FOUND_IN_SOURCES" else ""))),
    ("Website status", lambda r: r["website_status"]),
    ("Website note", lambda r: "; ".join(x for x in (r["website_note"], r["website_search_note"]) if x)),
    ("Google Business Profile", lambda r: r["gbp_url"] or r["gbp_status"]),
    ("GBP website field", lambda r: r["gbp_website"]),
    ("GBP reviews", lambda r: r["gbp_reviews"]),
    ("Facebook URL", lambda r: r["facebook_url"]),
    ("Facebook last post", lambda r: r["facebook_last_post"] or (r["facebook_status"] if r["facebook_url"] else "")),
    ("Instagram URL", lambda r: r["instagram_url"]),
    ("Instagram last post", lambda r: r["instagram_last_post"] or (r["instagram_status"] if r["instagram_url"] else "")),
    ("TikTok URL", lambda r: r["tiktok_url"]),
    ("TikTok last post", lambda r: r["tiktok_last_post"] or (r["tiktok_status"] if r["tiktok_url"] else "")),
    ("Last activity (all platforms)", lambda r: r["last_activity_date"]),
    ("Days since last activity", lambda r: int(r["days_since_activity"]) if str(r["days_since_activity"]).isdigit() else ""),
    ("Activity status", lambda r: r["activity_status"]),
    ("Meta Ads status", lambda r: r["meta_ads_status"] or "NOT CHECKED"),
    ("Active ads (count)", lambda r: r["meta_ads_count"]),
    ("Ad Library URL", lambda r: r["ad_library_url"]),
    ("Public business email", lambda r: r["email"]),
    ("Email source", lambda r: r["email_source"]),
    ("Public business phone", lambda r: r["phone"]),
    ("Phone source", lambda r: r["phone_source"]),
    ("Score explanation", lambda r: r["score_breakdown"]),
    ("Notes", lambda r: " | ".join(x for x in (
        f"CONFLICT: {r['conflicts']}" if r["conflicts"] else "", r["notes"], r["caen_note"],
        f"identity: {r['identity_evidence']}" if r["identity_evidence"] else "", r["social_search_note"]) if x)),
    ("Source URLs", lambda r: r["sources"]),
    ("Checked at", lambda r: r["checked_at"]),
]


def _score(r):
    return int(r["score"]) if str(r["score"]).isdigit() else -1


def leads_with(rows, cat=None):
    sel = [r for r in rows if r["categories"] and (cat is None or cat in [x.strip() for x in r["categories"].split(",")])]
    # Belt and braces: deduplicate by CUI, then by name + city (the register already merges on insert).
    seen, out = set(), []
    for r in sorted(sel, key=_score, reverse=True):
        key = ("cui", C.norm_cui(r["cui"])) if C.norm_cui(r["cui"]) else (
            "nc", C.norm_name(r["brand_name"] or r["legal_name"]), C.city_key(r["city"]))
        if key not in seen:
            seen.add(key)
            out.append(r)
    return out


def needs_review(rows):
    out = []
    for r in rows:
        if r["categories"] or r["verification_status"] == "DISCARDED":
            continue
        why = []
        if r["verification_status"] in ("CONFLICT", "UNVERIFIED") and "cat1" in r["origin"]:
            why.append(f"registry verification {r['verification_status']}")
        if r["activity_status"] == "UNKNOWN":
            why.append("activity UNKNOWN")
        if r["activity_status"] == "INACTIVE" and (r["meta_ads_status"] or "NOT CHECKED") == "NOT CHECKED":
            why.append("inactive, ads not checked")
        if r["website_status"] in ("NOT_FOUND_IN_SOURCES", "UNKNOWN"):
            why.append(f"website {r['website_status']}")
        if r["gbp_status"] == "AMBIGUOUS":
            why.append("GBP ambiguous")
        if why:
            r = dict(r)
            r["score_breakdown"] = "; ".join(why) + (f" | {r['score_breakdown']}" if r["score_breakdown"] else "")
            out.append(r)
    return out


def summary_lines(rows, log_rows):
    lines = [f"Lead research run — {C.now().strftime('%Y-%m-%d %H:%M')} ({C.cfg()['timezone']})", ""]
    lines.append("Leads per category:")
    for name, cat in TABS:
        lines.append(f"  {name}: {len(leads_with(rows, cat))}")
    allr = leads_with(rows)
    pri = {p: sum(1 for r in allr if r["priority"] == p) for p in ("HIGH", "MEDIUM", "LOW")}
    lines.append(f"  All leads (deduplicated): {len(allr)}  — HIGH {pri['HIGH']}, MEDIUM {pri['MEDIUM']}, LOW {pri['LOW']}")
    lines += ["", "Top 10 leads:"]
    for i, r in enumerate(allr[:10], 1):
        lines.append(f"  {i:2}. [{r['score']} {r['priority']}] {r['brand_name'] or r['legal_name']} — {r['city']} — "
                     f"cat {r['categories']} — {r['score_breakdown'][:110]}")
    live = [r for r in rows if r["verification_status"] != "DISCARDED"]
    issues = {
        "registry conflicts": sum(1 for r in live if r["conflicts"]),
        "Cat 1 candidates not confirmed by a 2nd source": sum(
            1 for r in live if "cat1" in r["origin"] and r["verification_status"] != "VERIFIED"),
        "discarded (not found / not active / out of scope)": sum(1 for r in rows if r["verification_status"] == "DISCARDED"),
        "activity UNKNOWN": sum(1 for r in live if r["activity_status"] == "UNKNOWN"),
        "inactive but Meta ads NOT CHECKED": sum(1 for r in live if r["activity_status"] == "INACTIVE"
                                                 and (r["meta_ads_status"] or "NOT CHECKED") == "NOT CHECKED"),
        "website not found in sources (web search pending)": sum(1 for r in live if r["website_status"] == "NOT_FOUND_IN_SOURCES"),
        "website UNKNOWN (blocked/robots)": sum(1 for r in live if r["website_status"] == "UNKNOWN"),
        "GBP ambiguous": sum(1 for r in live if r["gbp_status"] == "AMBIGUOUS"),
        "no public contact": sum(1 for r in allr if not (r["email"] or r["phone"])),
    }
    lines += ["", "Data quality issues:"]
    lines += [f"  {k}: {v}" for k, v in issues.items() if v]
    blocked = sorted({f"{l['source']} ({l['blocked']})" for l in log_rows if l.get("blocked")})
    limits = [l["note"] for l in log_rows if "LIMITATION" in (l.get("note") or "")]
    errors = [f"{l['source']}: {l['errors']}" for l in log_rows if l.get("errors")]
    if blocked or limits or errors:
        lines += ["", "Source limitations this session:"]
        lines += [f"  blocked/not used: {b}" for b in blocked]
        lines += [f"  {x}" for x in dict.fromkeys(limits)]
        lines += [f"  error: {e}" for e in dict.fromkeys(errors)]
    return lines


def export(path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = C.load()
    log_rows = C.load_log(C.session_id())
    stamp = C.now().strftime("%Y%m%d-%H%M")
    path = path or os.path.join(C.data_dir("output"), f"leads_{stamp}.xlsx")
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    pri_fill = {"HIGH": "C6EFCE", "MEDIUM": "FFEB9C", "LOW": "F2F2F2"}

    def sheet(title, data, cols=OUT, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append([h for h, _ in cols])
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in data:
            ws.append([f(r) for _, f in cols])
            p = r.get("priority")
            if p in pri_fill:
                ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=pri_fill[p])
        ws.freeze_panes = "C2" if cols is OUT else "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, (h, _) in enumerate(cols, 1):
            width = 14 if cols is OUT else 22
            if h in ("Legal name", "Brand name", "Website URL", "Notes", "Score explanation", "Source URLs",
                     "note", "query", "Facebook URL", "Instagram URL", "Ad Library URL", "Website note"):
                width = 40
            ws.column_dimensions[get_column_letter(i)].width = width
        return ws

    lines = summary_lines(rows, log_rows)
    ws = wb.active
    ws.title = "Summary"
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 160
    ws["A1"].font = Font(bold=True, size=13)
    for name, cat in TABS:
        sheet(name, leads_with(rows, cat))
    sheet(ALL_TAB, leads_with(rows))
    sheet("Needs Review", needs_review(rows))
    log_cols = [(h, (lambda k: lambda r: r.get(k, ""))(h)) for h in C.LOG_COLUMNS]
    sheet("Run Log", log_rows, cols=log_cols)
    wb.save(path)
    csv_path = path[:-5] + "_all_leads.csv"
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow([h for h, _ in OUT])
        for r in leads_with(rows):
            w.writerow([f(r) for _, f in OUT])
    print("\n".join(lines))
    print(f"\nExcel: {path}\nCSV:   {csv_path}")
    return path
